import csv
import requests
from bs4 import BeautifulSoup
import docx

def get_description(url):
    try:
        response = requests.get(url)
        soup = BeautifulSoup(response.text, 'html.parser')
        description = soup.find('p', class_='award-description')
        if description:
            return description.text.strip()
        else:
            return 'No description found'
    except Exception as e:
        return f'Error: {e}'

def main():
    # Try opening the cleaned CSV file first, if it fails, try the original one
    try:
        csv_file = open('scholarship_links_cleaned.csv', 'r')
    except FileNotFoundError:
        csv_file = open('scholarship_links.csv', 'r')

    reader = csv.reader(csv_file)
    # Create a new Word document
    doc = docx.Document()

    # Get the heading and body styles from the document
    heading_style = doc.styles['Heading 1']
    body_style = doc.styles['Body Text']

    # Iterate over each row in the CSV file
    for row in reader:
        url = row[0]  # Assuming the URL is in the first column
        award = row[1]  # Assuming the award is in the second column
        deadline = row[2]  # Assuming the deadline is in the third column
        description = get_description(url)

        # Add the URL and description to the Word document
        if url.strip().lower() != "view link":
            doc.add_paragraph(f'{url}', style=heading_style)
        if award.strip().lower() != "award":
            doc.add_paragraph(f'Award: {award}', style=body_style)
        if deadline.strip().lower() != "deadline":
            doc.add_paragraph(f'Deadline: {deadline}', style=body_style)

        if description.strip().lower() != "no description found":
            doc.add_paragraph(description, style=body_style)

    # Save the Word document
    doc.save('scholarship_descriptions.docx')

    # Close the CSV file
    csv_file.close()

if __name__ == '__main__':
    main()